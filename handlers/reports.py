# handlers/reports.py
from aiogram import Router, types, F
from aiogram.types import ReplyKeyboardMarkup, KeyboardButton, FSInputFile
from aiogram.fsm.context import FSMContext
from keyboards.main_menu import get_main_menu
from database.db_manager import get_db  # Контекстный менеджер для работы с БД
from database.models import Order, OrderStatus, LEDModule, Supplier, Customer
from constants import EXPORT_PATHS
from sqlalchemy import func
from datetime import datetime
import openpyxl
import os
import logging

logger = logging.getLogger(__name__)

router = Router()

# --- Клавиатура ---
def get_reports_keyboard():
    keyboard = [
        [KeyboardButton(text="📈 Общая сводка")],
        [KeyboardButton(text="📦 Отчет по заказам")],
        [KeyboardButton(text="🔲 Остатки LED модулей")],
        [KeyboardButton(text="🏆 Топ контрагентов")],
        [KeyboardButton(text="📄 Экспорт в Excel")],
        [KeyboardButton(text="⬅️ Назад")]
    ]
    return ReplyKeyboardMarkup(keyboard=keyboard, resize_keyboard=True)


# --- Главное меню отчётов ---
@router.message(F.text == "📊 Отчеты")
async def menu_reports(message: types.Message):
    await message.answer(
        "📊 <b>Отчеты и аналитика</b>\n\nВыберите тип отчёта:",
        reply_markup=get_reports_keyboard(),
        parse_mode="HTML"
    )


# --- Общая сводка ---
@router.message(F.text == "📈 Общая сводка")
async def report_summary(message: types.Message):
    try:
        with get_db() as db:
            total_orders = db.query(func.count(Order.id)).scalar()
            active_orders = db.query(func.count(Order.id)).filter(
                Order.status.in_([OrderStatus.CREATED, OrderStatus.IN_PROGRESS])
            ).scalar()
            total_modules = db.query(func.count(LEDModule.id)).scalar()
            total_qty = db.query(func.sum(LEDModule.quantity)).filter(
                LEDModule.status == 'in_stock'
            ).scalar() or 0
            total_suppliers = db.query(func.count(Supplier.id)).scalar()
            total_customers = db.query(func.count(Customer.id)).scalar()

        text = (
            f"📈 <b>Общая сводка</b>\n\n"
            f"📦 Заказы: <b>{total_orders}</b> (активные: {active_orders})\n"
            f"🔲 Модули: <b>{total_modules}</b> поз. ({int(total_qty)} шт)\n"
            f"🤝 Поставщиков: <b>{total_suppliers}</b>\n"
            f"👥 Покупателей: <b>{total_customers}</b>\n"
            f"📅 {datetime.now().strftime('%d.%m.%Y %H:%M')}"
        )
        await message.answer(text, parse_mode="HTML")
    except Exception as e:
        logger.error(f"Ошибка в отчёте 'Общая сводка': {e}", exc_info=True)
        await message.answer("❌ Не удалось получить сводку.")


# --- Отчёт по заказам ---
@router.message(F.text == "📦 Отчет по заказам")
async def report_orders(message: types.Message):
    try:
        with get_db() as db:
            stats = db.query(Order.status, func.count(Order.id)).group_by(Order.status).all()
            today = datetime.now().replace(hour=0, minute=0, second=0)
            orders_today = db.query(func.count(Order.id)).filter(Order.created_at >= today).scalar()

        status_names = {
            'created': '📝 Создан',
            'in_progress': '⚙️ В работе',
            'shipped': '🚚 Отгружен',
            'completed': '✅ Завершен',
            'cancelled': '❌ Отменен'
        }

        text = f"📦 <b>Отчет по заказам</b>\n\nЗа сегодня: <b>{orders_today}</b>\n\n"
        for status, count in stats:
            s_val = status.value if hasattr(status, 'value') else str(status)
            name = status_names.get(s_val, s_val)
            text += f"{name}: {count}\n"
        if not stats:
            text += "Нет данных"

        await message.answer(text, parse_mode="HTML")
    except Exception as e:
        logger.error(f"Ошибка в отчёте 'По заказам': {e}", exc_info=True)
        await message.answer("❌ Не удалось получить статистику по заказам.")


# --- Остатки LED модулей ---
@router.message(F.text == "🔲 Остатки LED модулей")
async def report_led_stock(message: types.Message):
    try:
        with get_db() as db:
            stock = db.query(
                LEDModule.step,
                func.count(LEDModule.id),
                func.sum(LEDModule.quantity)
            ).filter(LEDModule.status == 'in_stock') \
             .group_by(LEDModule.step) \
             .order_by(LEDModule.step).all()

        if not stock:
            await message.answer("❌ Склад пуст")
            return

        text = "🔲 <b>Остатки LED</b>\n\n"
        total_p, total_q = 0, 0
        for step, count, qty in stock:
            q = qty or 0
            text += f"▫️ <b>{step}</b>: {count} поз. ({int(q)} шт)\n"
            total_p += count
            total_q += q
        text += f"\n<b>Итого</b>: {total_p} поз. ({int(total_q)} шт)"

        await message.answer(text, parse_mode="HTML")
    except Exception as e:
        logger.error(f"Ошибка в отчёте 'Остатки LED': {e}", exc_info=True)
        await message.answer("❌ Не удалось получить остатки.")


# --- Топ контрагентов ---
@router.message(F.text == "🏆 Топ контрагентов")
async def report_top_partners(message: types.Message):
    try:
        with get_db() as db:
            top_cust = db.query(Customer.name, func.count(Order.id)) \
                .join(Order) \
                .group_by(Customer.id) \
                .order_by(func.count(Order.id).desc()) \
                .limit(5).all()

            top_supp = db.query(Supplier.name, func.count(Order.id)) \
                .join(Order) \
                .group_by(Supplier.id) \
                .order_by(func.count(Order.id).desc()) \
                .limit(5).all()

        text = "🏆 <b>Топ контрагентов</b>\n\n"

        text += "<b>Покупатели:</b>\n"
        if top_cust:
            for i, (name, count) in enumerate(top_cust, 1):
                text += f"{i}. {name} — {count}\n"
        else:
            text += "Нет данных\n"

        text += "\n<b>Поставщики:</b>\n"
        if top_supp:
            for i, (name, count) in enumerate(top_supp, 1):
                text += f"{i}. {name} — {count}\n"
        else:
            text += "Нет данных"

        await message.answer(text, parse_mode="HTML")
    except Exception as e:
        logger.error(f"Ошибка в отчёте 'Топ контрагентов': {e}", exc_info=True)
        await message.answer("❌ Не удалось получить топ контрагентов.")


# --- Экспорт в Excel ---
@router.message(F.text == "📄 Экспорт в Excel")
async def export_to_excel(message: types.Message):
    try:
        await message.answer("⏳ Формируем Excel-отчёт...")

        with get_db() as db:
            # Лимиты для безопасности
            orders = db.query(Order).order_by(Order.created_at.desc()).limit(1000).all()
            modules = db.query(LEDModule).filter(LEDModule.status == 'in_stock').limit(5000).all()
            suppliers = db.query(Supplier).limit(1000).all()
            customers = db.query(Customer).limit(1000).all()

        wb = openpyxl.Workbook()

        # === Лист: Заказы ===
        ws = wb.active
        ws.title = "Заказы"
        ws.append(["ID", "Номер", "Поставщик", "Покупатель", "Статус", "Создан", "Отгружен"])
        for order in orders:
            ws.append([
                order.id,
                order.order_number,
                order.supplier.name if order.supplier else "Н/Д",
                order.customer.name if order.customer else "Н/Д",
                order.status.value,
                order.created_at,
                order.ship_date
            ])
        # Формат дат
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=6, max_col=7):
            for cell in row:
                if cell.value:
                    cell.number_format = "DD.MM.YYYY HH:MM"

        # === Лист: LED Модули ===
        ws = wb.create_sheet("LED Модули")
        ws.append(["ID", "Шаг", "Модель", "Кол-во", "Ширина", "Высота", "Место", "Статус", "Получен"])
        for m in modules:
            ws.append([
                m.id, m.step, m.model or "", m.quantity,
                m.width or "", m.height or "", m.location or "",
                m.status.value, m.received_date
            ])
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=9, max_col=9):
            for cell in row:
                if cell.value:
                    cell.number_format = "DD.MM.YYYY"

        # === Лист: Поставщики ===
        ws = wb.create_sheet("Поставщики")
        ws.append(["ID", "Название", "Контакт", "Телефон", "Email", "Адрес"])
        for s in suppliers:
            ws.append([
                s.id, s.name, s.contact_person or "", s.phone or "",
                s.email or "", s.address or ""
            ])

        # === Лист: Покупатели ===
        ws = wb.create_sheet("Покупатели")
        ws.append(["ID", "Название", "Контакт", "Телефон", "Email", "Адрес"])
        for c in customers:
            ws.append([
                c.id, c.name, c.contact_person or "", c.phone or "",
                c.email or "", c.address or ""
            ])

        # Сохранение
        export_dir = EXPORT_PATHS["reports_dir"]
        os.makedirs(export_dir, exist_ok=True)
        filename = f"{export_dir}/warehouse_{datetime.now().strftime('%Y%m%d_%H%M%S')}{EXPORT_PATHS['excel_extension']}"
        wb.save(filename)

        # Отправка
        await message.answer_document(
            document=FSInputFile(filename),
            caption=f"📄 <b>Отчёт сформирован</b>\n\n"
                   f"📊 Данных экспортировано:\n"
                   f" • Заказы: {len(orders)}\n"
                   f" • LED модули: {len(modules)}\n"
                   f" • Поставщики: {len(suppliers)}\n"
                   f" • Покупатели: {len(customers)}\n\n"
                   f"⏰ {datetime.now().strftime('%d.%m.%Y %H:%M')}",
            parse_mode="HTML"
        )

        # Удаление файла
        os.remove(filename)
    except Exception as e:
        logger.error(f"Ошибка экспорта в Excel: {e}", exc_info=True)
        await message.answer("❌ Не удалось экспортировать данные.")
    finally:
        if 'filename' in locals() and os.path.exists(filename):
            os.remove(filename)


# --- Назад в главное меню ---
@router.message(F.text == "⬅️ Назад")
async def menu_back(message: types.Message, state: FSMContext):
    await state.clear()  # Очищаем состояние
    await message.answer("🏠 Главное меню", reply_markup=get_main_menu())